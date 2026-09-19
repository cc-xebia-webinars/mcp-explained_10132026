"""Build the webinar slide deck from the Xebia template.

This script is the only way the deck changes. The PPTX in `docs/` is generated from
the SLIDES list below and should never be edited by hand: change the data here,
rebuild, regenerate the PDF, and then update the matching heading in
PRESENTER_SCRIPT.md.

Run from the repository root:

    uv run --with python-pptx python instructor/build_deck.py
    uv run --with python-pptx python instructor/build_deck.py --inventory

Then convert to PDF with LibreOffice, as described in the root CLAUDE.md.
"""

from __future__ import annotations

import argparse
import copy
import io
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "instructor" / "Xebia PPT General Template.pptx"
OUTPUT = ROOT / "docs" / "mcp-explained-webinar-slides.pptx"

# Diagrams are drawn with the diagram-builder MCP server. Each .svg in this folder is
# the editable source, and the .png next to it is the export the deck embeds.
DIAGRAMS = ROOT / "instructor" / "diagrams"

# Layout indexes in the Xebia template, named so the slide list below reads clearly.
COVER = 0          # Main-cover_light
BASIC = 6          # Content_Basic
THREE_COL = 8      # Content_3 Columns_Comparison_2
TWO_COL = 9        # Content_2 Columns
BREAK = 24         # 1_Photo_background
END = 27           # End-cover_light

# The template's own body default is 12 pt, which is unreadable on a shared screen,
# so every run gets an explicit size instead of inheriting one.
BULLET_PT = 20
CODE_PT = 18
COLUMN_CODE_PT = 16
COLUMN_BODY_PT = 16
COLUMN_HEADING_PT = 16
SUBTITLE_PT = 16
MONO = "Consolas"

# Density rules from CLAUDE.md, checked against this data before anything is built.
MAX_SLIDES = 20
MAX_BULLETS = 6
MAX_WORDS = 11     # "under a dozen words"
MAX_CODE_LINES = 12

# Column placeholders aren't in left-to-right index order, so each column is named by
# its (heading idx, body idx) pair, measured from the template.
THREE_COL_SLOTS = [(11, 1), (17, 16), (14, 13)]    # left, middle, right
TWO_COL_SLOTS = [(11, 1), (14, 13)]                # left, right

# The two-column layout reserves a band for icon pictures above each column. We don't
# use icons, and once those frames are removed the columns sit alone in the bottom half
# of the slide. Moving them up to the three-column layout's positions closes the gap.
TWO_COL_HEADING_TOP_IN = 2.16
TWO_COL_BODY_TOP_IN = 2.84
TWO_COL_BODY_HEIGHT_IN = 4.06

# Diagram slides fill the body area of the basic layout, under the title.
DIAGRAM_LEFT_IN = 0.61
DIAGRAM_TOP_IN = 1.85
DIAGRAM_MAX_HEIGHT_IN = 5.0
DIAGRAM_MAX_WIDTH_IN = 12.1


# ---------------------------------------------------------------------------- slides
#
# Text wrapped in backticks is set in a monospace font, so method names, file paths
# and code read as code. Everything else is set in the template's Arial.

SLIDES = [
    {
        "kind": "Title",
        "layout": COVER,
        # The line break keeps the title clear of the cover's artwork.
        "title": "Model Context Protocol\n(MCP) Explained",
        "subtitle": ("Connecting AI Agents to Your Tools and Data\n"
                     "Eric Greene · Xebia\neric.greene@xebia.com"),
        "notes": (
            "Welcome and introductions. Ninety minutes: about twenty of slides, seventy of "
            "live demos in VS Code with GitHub Copilot. Questions are welcome at any time, in "
            "the chat or out loud, and the demos leave plenty of room to answer them. "
            "Everything is in the webinar repo. Mention early that the MCP spec moved to a "
            "stateless revision, 2026-07-28, and the Python SDK to version 2, so older "
            "tutorials may look different."
        ),
    },
    {
        "kind": "Agenda",
        "layout": BASIC,
        "title": "Agenda",
        "bullets": [
            "What MCP is, and why it exists",
            "Hosts, clients, servers, and the lifecycle",
            "Tools, resources, and prompts",
            "Ready-made servers you can use today",
            "Building a custom server in Python",
            "Security, trust boundaries, and production audits",
        ],
        "notes": (
            "Six topics, matching the README agenda. Topics one, two and four lead into Demo "
            "1. Topics three and five lead into Demos 2 and 3. Topic six leads into Demo 4. "
            "Keep this brief; the slides set up each demo rather than replacing it."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "What MCP is",
        "bullets": [
            "An open protocol connecting AI apps to tools and data",
            "JSON-RPC 2.0 messages over stdio or HTTP",
            "Servers describe their capabilities; hosts discover them at runtime",
            "Supported across VS Code, Claude, Cursor, and more",
            "A versioned specification; `2026-07-28` is current",
            "Builds on function calling rather than replacing it",
        ],
        "notes": (
            "MCP is a standard way for an AI application to find and use tools and data it "
            "wasn't built with. Messages are JSON-RPC 2.0. The key idea is runtime discovery: "
            "the host asks the server what it offers instead of having it hard-coded. The "
            "current revision is 2026-07-28, which made the protocol stateless. Under the "
            "hood the model still uses function calling; MCP standardizes where the functions "
            "come from."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "How MCP compares",
        "columns": [
            ("Ad-hoc integrations", [
                "Every app wires up every tool itself",
                "N apps and M tools means N×M connectors",
            ]),
            ("Function calling", [
                "Tools defined per model API",
                "Schemas and code live inside each app",
            ]),
            ("MCP", [
                "Build a server once, use it anywhere",
                "Any MCP host discovers it at runtime",
            ]),
        ],
        "notes": (
            "Left: custom glue code per app per tool, which doesn't scale. Middle: function "
            "calling lets a model request a function, but each app defines and implements its "
            "functions for one model API. Right: MCP moves the tool into a server that any "
            "host can use, so N times M becomes N plus M. Be fair: for one app with two "
            "private functions, plain function calling is fine."
        ),
    },
    {
        "kind": "Diagram",
        "layout": BASIC,
        "title": "Hosts, clients, and servers",
        "image": "architecture.png",
        # Trims the diagram page's empty top and bottom margins, as fractions.
        "crop": (0.05, 0.10),
        "notes": (
            "Walk the diagram left to right. The host is the application, and it owns the "
            "model and the user's trust. Inside it, one client per server holds the "
            "connection. Servers expose capabilities and know nothing about the model or "
            "about each other. That separation is why one "
            "host can combine many servers, and why the host is where approvals belong."
        ),
    },
    {
        "kind": "Diagram",
        "layout": BASIC,
        "title": "The request/response lifecycle",
        "image": "lifecycle.png",
        "crop": (0.05, 0.06),
        "notes": (
            "Left lane, modern servers (2026-07-28), are stateless: every request carries the protocol "
            "version and client details in _meta, and server/discover replaces the handshake. "
            "Right lane, legacy servers (2025-11-25 and earlier), still expect initialize first, and many "
            "popular servers are legacy today, including the official filesystem server. "
            "Clients probe with server/discover and fall back. Demo 1 shows both on the wire."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Ready-made servers",
        "columns": [
            ("Local, over stdio", [
                "Filesystem: read and search project files",
                "Databases: SQLite, Postgres, and others",
                "Git: history, diffs, and branches",
            ]),
            ("Hosted, over HTTP", [
                "GitHub: issues, pull requests, and code",
                "Microsoft Learn: current official documentation",
                "Search and SaaS servers from vendors",
            ]),
        ],
        "notes": (
            "Before building anything, check whether a server exists. Local servers run as a "
            "child process and see your machine. Hosted servers run elsewhere and sign in with "
            "OAuth. Demo 1 uses the filesystem server, GitHub's hosted server in read-only "
            "mode, and Microsoft Learn as the documentation search example. Database servers "
            "exist too; our database example is the custom server in Demo 2."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 1: Meet MCP",
        "subtitle": "demos/01-meet-mcp · about 16 minutes",
        "notes": (
            "Run lifecycle.py --step against the filesystem server (legacy handshake), then "
            "against hello_server.py (modern, stateless), pressing Enter phase by phase. Then .vscode/mcp.json, start the three "
            "servers, and use filesystem, GitHub (read-only) and Microsoft Learn, finishing "
            "with one prompt that uses two servers. Point out the roots surprise on the "
            "filesystem server."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Tools, resources, and prompts",
        "columns": [
            ("Tools", [
                "Actions the model decides to call",
                "`search_services`, `create_incident`",
            ]),
            ("Resources", [
                "Context the user or host attaches",
                "`catalog://teams`, service runbooks",
            ]),
            ("Prompts", [
                "Templates the user chooses to run",
                "`incident_briefing`, `team_onboarding`",
            ]),
        ],
        "notes": (
            "The difference is who's in control. The model chooses tools. The user or host "
            "chooses resources, which are read-only context identified by a URI. The user "
            "chooses prompts, which are reusable instructions with arguments. The examples are "
            "the real names from Demo 2's service catalog server."
        ),
    },
    {
        "kind": "Content",
        "layout": TWO_COL,
        "title": "Choosing the right primitive",
        "columns": [
            ("Make it a tool when", [
                "The model should decide when it's needed",
                "It takes arguments or changes something",
                "The answer depends on the question",
            ]),
            ("Make it a resource or prompt when", [
                "The user knows which context matters",
                "It's a document, not an action",
                "A team wants one repeatable workflow",
            ]),
        ],
        "notes": (
            "Most servers start and end with tools, because every host supports them. Use "
            "resources for documents the user wants to attach deliberately, like a runbook. "
            "Use prompts to package a team workflow so everyone gets the same structure. If "
            "you're unsure, start with a tool."
        ),
    },
    {
        "kind": "Code",
        "layout": BASIC,
        "title": "A tool in a few lines of Python",
        "code": [
            "from mcp.server import MCPServer",
            "",
            'mcp = MCPServer("service-catalog")',
            "",
            "@mcp.tool(annotations=READ_ONLY)",
            "def get_service(name: str) -> Service:",
            '    """Owner, on-call, and dependencies."""',
            "    return catalog_db.get_service(name)",
            "",
            'mcp.run(transport="stdio")',
        ],
        "notes": (
            "This is the official Python SDK, version 2. MCPServer replaced FastMCP, so older "
            "tutorials use a different import. The function name becomes the tool name, the "
            "type hints become the input schema, the return type becomes the output schema, "
            "and the docstring is what the model reads. Annotations mark it read-only. The run "
            "call picks the transport, and that's the only line Demo 3 changes."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 2: A custom server over stdio",
        "subtitle": "demos/02-custom-server-stdio · about 18 minutes",
        "notes": (
            "Walk catalog_db.py and server.py, run check_server.py, then the Inspector CLI. "
            "Connect it in Copilot and ask who's on call for the payments API, attach the "
            "catalog://teams resource, run the incident_briefing prompt, and open an incident "
            "to show the approval. Reset with uv run catalog_db.py --reset."
        ),
    },
    {
        "kind": "Code",
        "layout": TWO_COL,
        "title": "stdio or Streamable HTTP",
        "code_columns": [
            ("stdio: the host starts the server", [
                "# server.py",
                'mcp.run(transport="stdio")',
                "",
                "",  # keeps the two mcp.json blocks level with each other
                "# .vscode/mcp.json",
                '"type": "stdio",',
                '"command": "uv",',
                '"args": ["run", "server.py"]',
            ]),
            ("Streamable HTTP: the server runs at a URL", [
                "# http_server.py",
                'mcp.run(transport="streamable-http",',
                '        host="127.0.0.1", port=8765)',
                "",
                "# .vscode/mcp.json",
                '"type": "http",',
                '"url": "http://127.0.0.1:8765/mcp"',
            ]),
        ],
        "notes": (
            "Same server, two transports, and this is everything that changes: the run call "
            "and the host configuration. stdio suits personal, local tools: the host starts "
            "the process, and there's no network. Streamable HTTP suits shared services: one "
            "endpoint, each message an HTTP POST, many clients. Since 2026-07-28 there are no "
            "sessions, so HTTP servers scale like any stateless web service. "
            "HTTP brings web risks with it, and Demo 3 shows two guards: Host checks and a "
            "bearer token."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 3: The same server over HTTP",
        "subtitle": "demos/03-streamable-http · about 16 minutes",
        "notes": (
            "Show server.py is identical with git diff --no-index. Walk http_server.py, start "
            "it, point Copilot at the URL. Run http_client.py as a second client and show "
            "Copilot sees its incident. Run rebinding_check.py: 200 for the honest request, "
            "421 for the forged Host header. Finish with secure_http_server.py and "
            "auth_check.py, then switch Copilot to the token-protected entry."
        ),
    },
    {
        "kind": "Content",
        "layout": THREE_COL,
        "title": "Where the trust boundaries are",
        "columns": [
            ("Model", [
                "Reads every tool description as guidance",
                "Can be steered by injected text",
            ]),
            ("Host", [
                "Approves calls and shows their arguments",
                "Decides which servers are connected",
            ]),
            ("Server", [
                "Runs code with your permissions",
                "Reaches your files, data, and networks",
            ]),
        ],
        "notes": (
            "Every server can put text in front of the model and run code as you. The model "
            "can't reliably tell instructions from data, so anything it reads, including tool "
            "descriptions and tool results, can steer it. The host is the control point: "
            "which servers are connected, and which calls need approval."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "Risks to plan for",
        "bullets": [
            "Tool poisoning: instructions hidden in tool descriptions",
            "Prompt injection through documents, issues, and web pages",
            "Over-broad access: wide roots, tokens, and write tools",
            "Raw query tools that run whatever the model writes",
            "Server updates that change tools after approval",
            "HTTP servers exposed without auth or Host checks",
        ],
        "notes": (
            "Tool poisoning is in the server itself. Prompt injection arrives in data a tool "
            "returns, such as an issue body. Over-broad access turns either one into real "
            "damage. Raw SQL or shell tools hand the model a blank cheque. A server can change "
            "its descriptions after you approve it. And an HTTP server is a web server, with "
            "all that implies. Demo 3 shows the last one, and Demo 4 shows the other five: "
            "poisoning arriving in an update, injection through data, over-broad access and "
            "raw SQL."
        ),
    },
    {
        "kind": "Content",
        "layout": BASIC,
        "title": "Audit before production",
        "bullets": [
            "Read every tool description, and pin server versions",
            "Grant the narrowest root, scopes, and toolsets",
            "Prefer read-only endpoints and parameterized tools",
            "Keep approvals on for anything that writes",
            "Require OAuth, and bind HTTP servers narrowly",
            "Log tool calls so you can review them",
        ],
        "notes": (
            "Treat an MCP server like a dependency with production access. Review what it "
            "tells the model, pin it, and recheck on upgrade. Least privilege everywhere: "
            "roots, token scopes, GitHub toolsets and read-only mode. Specific tools beat "
            "general ones. Approvals only help if people read the arguments. For shared "
            "servers, OAuth and a narrow network exposure are the baseline. Demo 4 ends on a "
            "tool-call log. The full list is in docs/mcp-review-checklist.md."
        ),
    },
    {
        "kind": "Demo",
        "layout": BREAK,
        "title": "Demo 4: Security and trust boundaries",
        "subtitle": "demos/04-security · about 20 minutes",
        "notes": (
            "Start all three servers and run audit_tools.py --approve. Change "
            "CONVERTER_RELEASE to 1.1, restart it, ask to convert 42 km to miles and watch for "
            "the poisoned read. Rerun audit_tools.py to catch the change. List files with the wide root, then scope it to "
            "shared-docs and show the traversal refused. Delete resolved incidents through "
            "run_sql, then try injection on incidents_for_service. Summarize checkout-web's incidents "
            "to show injection through data, then open data/tool-calls.log. Reset at the end."
        ),
    },
    {
        "kind": "Wrap",
        "layout": BASIC,
        "title": "Next steps",
        "bullets": [
            "Clone the repo and run the four demos",
            "Connect one ready-made server to your daily work",
            "Wrap one internal API as a small MCP server",
            "Adapt the repo's review checklist for your team",
            "Follow the specification changelog; it moves quickly",
        ],
        "notes": (
            "Point people to demos/README.md for setup. Suggest starting with a read-only "
            "server, then wrapping one internal API they already know. The review checklist "
            "is linked from the repo README, next to the slides. Mention the Xebia Academy "
            "link there too."
        ),
    },
    {
        "kind": "Closing",
        "layout": END,
        "title": "Thank you",
        "subtitle": "Eric Greene · Xebia\neric.greene@xebia.com",
        "notes": (
            "Thank the audience. Leave this slide up while you answer any final questions, and "
            "repeat that the slides and demos are in the repo."
        ),
    },
]


# ---------------------------------------------------------------------- validation

def words(text: str) -> int:
    """Count words the way a reader would: tokens containing a letter or digit."""
    return sum(1 for t in text.replace("`", "").split() if re.search(r"[A-Za-z0-9]", t))


def validate_data(slides: list[dict]) -> list[str]:
    problems = []
    if len(slides) > MAX_SLIDES:
        problems.append(f"{len(slides)} slides, limit is {MAX_SLIDES}")
    for n, s in enumerate(slides, start=1):
        groups = []
        if "bullets" in s:
            groups.append(s["bullets"])
        for _, items in s.get("columns", []):
            groups.append(items)
        total = sum(len(g) for g in groups)
        if total > MAX_BULLETS:
            problems.append(f"slide {n}: {total} bullets, limit is {MAX_BULLETS}")
        for g in groups:
            for b in g:
                if words(b) > MAX_WORDS:
                    problems.append(f"slide {n}: {words(b)} words in {b!r}")
        code_blocks = [s.get("code", [])] + [lines for _, lines in s.get("code_columns", [])]
        for block in code_blocks:
            if len(block) > MAX_CODE_LINES:
                problems.append(f"slide {n}: {len(block)} code lines, limit is {MAX_CODE_LINES}")
        if "image" in s and not (DIAGRAMS / s["image"]).is_file():
            problems.append(f"slide {n}: diagram {s['image']} is missing; export it first")
        if not s.get("notes", "").strip():
            problems.append(f"slide {n}: speaker notes are empty")
    for n in (1, len(slides)):
        if "eric.greene@xebia.com" not in slides[n - 1].get("subtitle", ""):
            problems.append(f"slide {n}: the presenter's email must be on the first and last slides")
    return problems


# ---------------------------------------------------------------------------- build

def add_runs(paragraph, text: str, size: int) -> None:
    """Split on backticks so `code` segments get the monospace font."""
    for i, part in enumerate(text.split("`")):
        if not part:
            continue
        run = paragraph.add_run()
        run.text = part
        run.font.size = Pt(size)
        if i % 2 == 1:
            run.font.name = MONO


def fill_lines(placeholder, lines: list[str], size: int) -> None:
    """Write one paragraph per line, without leaving a stray empty first paragraph."""
    tf = placeholder.text_frame
    tf.clear()
    first = True
    for line in lines:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        add_runs(para, line, size)


def fill_code(placeholder, lines: list[str], size: int) -> None:
    """Write code as unbulleted monospace lines.

    The body placeholder bullets every paragraph, which would put a dot in front of
    each line of code, so each paragraph gets an explicit buNone and no indent.
    """
    tf = placeholder.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "0")
        p_pr.set("indent", "0")
        p_pr.insert(0, p_pr.makeelement(qn("a:buNone"), {}))
        run = para.add_run()
        # A non-breaking space keeps blank lines and leading indentation intact.
        run.text = line.replace(" ", " ") if line else " "
        run.font.size = Pt(size)
        run.font.name = MONO


def transparent_png(path: Path) -> io.BytesIO:
    """Make the diagram's white page background transparent.

    The diagram tool exports on solid white, which would cover the template's light
    background artwork. Pillow comes with python-pptx, so no extra dependency is needed.
    """
    img = Image.open(path).convert("RGBA")
    img.putdata([(r, g, b, 0) if min(r, g, b) >= 250 else (r, g, b, a)
                 for r, g, b, a in img.get_flattened_data()])
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def add_diagram(slide, image: str, crop: tuple[float, float]) -> None:
    """Place a diagram under the title, scaled to the body area and cropped.

    The diagram tool always exports a full landscape page, so the empty margins above
    and below the drawing are cropped off before the picture is sized.
    """
    top_crop, bottom_crop = crop
    pic = slide.shapes.add_picture(transparent_png(DIAGRAMS / image), Inches(DIAGRAM_LEFT_IN),
                                   Inches(DIAGRAM_TOP_IN))
    pic.crop_top, pic.crop_bottom = top_crop, bottom_crop
    aspect = pic.image.size[0] / (pic.image.size[1] * (1 - top_crop - bottom_crop))
    height = min(DIAGRAM_MAX_HEIGHT_IN, DIAGRAM_MAX_WIDTH_IN / aspect)
    pic.height, pic.width = Inches(height), Inches(height * aspect)


def move(placeholder, *, top: float, height: float | None = None) -> None:
    """Reposition a placeholder vertically, keeping its inherited left and width.

    A slide placeholder inherits its geometry from the layout until it's given its own,
    and python-pptx needs all four values set to write one, so the current left and
    width are carried over explicitly.
    """
    left, width = placeholder.left, placeholder.width
    new_height = Inches(height) if height is not None else placeholder.height
    placeholder.left, placeholder.width = left, width
    placeholder.top, placeholder.height = Inches(top), new_height


def restore_slide_number(slide) -> None:
    """python-pptx skips slide-number placeholders when cloning a layout, but the
    template's own slides carry them, so copy the layout's one across."""
    for lp in slide.slide_layout.placeholders:
        if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            slide.shapes._spTree.append(copy.deepcopy(lp._element))


def remove_unfilled(slide, filled: set[int]) -> None:
    """An unfilled picture placeholder renders as an empty frame, so anything we
    didn't write to is removed rather than left behind."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in filled:
            ph._element.getparent().remove(ph._element)


def build(output: Path) -> None:
    prs = Presentation(TEMPLATE)
    layouts = prs.slide_masters[0].slide_layouts
    assert len(layouts) == 29, f"template has {len(layouts)} layouts, expected 29"

    # Remove the template's example slides. Layouts, the master and the notes master
    # survive, which is all the deck needs.
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    id_list = prs.slides._sldIdLst
    for sld in list(id_list):
        prs.part.drop_rel(sld.get(rel_id))
        id_list.remove(sld)
    assert len(prs.slides) == 0

    for s in SLIDES:
        slide = prs.slides.add_slide(layouts[s["layout"]])
        ph = {p.placeholder_format.idx: p for p in slide.placeholders}
        filled = {0}
        ph[0].text_frame.text = s["title"]

        if s["layout"] in (COVER, BREAK, END):
            fill_lines(ph[1], s["subtitle"].split("\n"), SUBTITLE_PT)
            filled.add(1)
        elif "columns" in s:
            slots = THREE_COL_SLOTS if s["layout"] == THREE_COL else TWO_COL_SLOTS
            for (head_idx, body_idx), (heading, items) in zip(slots, s["columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_lines(ph[body_idx], items, COLUMN_BODY_PT)
                filled.update({head_idx, body_idx})
                if s["layout"] == TWO_COL:
                    move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                    move(ph[body_idx], top=TWO_COL_BODY_TOP_IN,
                         height=TWO_COL_BODY_HEIGHT_IN)
        elif "code_columns" in s:
            for (head_idx, body_idx), (heading, lines) in zip(TWO_COL_SLOTS, s["code_columns"]):
                fill_lines(ph[head_idx], [heading], COLUMN_HEADING_PT)
                fill_code(ph[body_idx], lines, COLUMN_CODE_PT)
                filled.update({head_idx, body_idx})
                move(ph[head_idx], top=TWO_COL_HEADING_TOP_IN)
                move(ph[body_idx], top=TWO_COL_BODY_TOP_IN, height=TWO_COL_BODY_HEIGHT_IN)
        elif "code" in s:
            fill_code(ph[1], s["code"], CODE_PT)
            filled.add(1)
        elif "image" in s:
            add_diagram(slide, s["image"], s["crop"])
        else:
            fill_lines(ph[1], s["bullets"], BULLET_PT)
            filled.add(1)

        remove_unfilled(slide, filled)
        if s["layout"] not in (COVER, END):
            restore_slide_number(slide)
        slide.notes_slide.notes_text_frame.text = s["notes"]

    assert len(prs.slides) == len(SLIDES)
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no notes"

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def inventory() -> str:
    rows = [f"{'#':>2}  {'kind':<8} title"]
    for n, s in enumerate(SLIDES, start=1):
        title = s["title"].replace("\n", " ")
        rows.append(f"{n:>2}  {s['kind']:<8} {title}")
    return "\n".join(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inventory", action="store_true",
                        help="print the slide list and exit, for checking the script against it")
    args = parser.parse_args()

    if args.inventory:
        print(inventory())
        return 0

    problems = validate_data(SLIDES)
    if problems:
        print("Slide data breaks the rules in CLAUDE.md:", file=sys.stderr)
        for p in problems:
            print(f"  - {p}", file=sys.stderr)
        return 1

    build(OUTPUT)
    print(f"built {OUTPUT.relative_to(ROOT)} with {len(SLIDES)} slides")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
